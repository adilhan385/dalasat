"""
create_word_manual.py
Скрипт генерации официального руководства по локальному запуску DalaSat в формате Microsoft Word (.docx).
Для экспертной комиссии хакатона «AgriTech AI Hackathon» (Aqmola Hub / Digital Aqmola).
"""

import os
DIR = os.path.dirname(os.path.abspath(__file__))
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    """Установка фонового цвета ячейки таблицы"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Установка внутренних отступов ячейки"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_manual():
    doc = Document()

    # Поля страницы: умеренные (2 см)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Цветовая палитра:
    C_DARK = RGBColor(0x1E, 0x25, 0x22)       # Основной графитовый
    C_GREEN = RGBColor(0x28, 0x54, 0x3E)      # Фирменный хвойный зеленый
    C_GRAY = RGBColor(0x5A, 0x65, 0x60)       # Второстепенный серый
    HEX_BEIGE = "F5F2EB"                      # Теплый бежевый
    HEX_LIGHT_GREEN = "EAEFEA"                # Нежный зеленый для плашек
    HEX_BORDER = "DDD7CB"

    # ─── 1. ТИТУЛЬНАЯ ПЛАШКА ──────────────────────────────────────────────────
    t_header = doc.add_table(rows=1, cols=1)
    t_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t_header.cell(0, 0)
    set_cell_background(c, HEX_BEIGE)
    set_cell_margins(c, top=200, bottom=200, left=250, right=250)

    p_badge = c.paragraphs[0]
    p_badge.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_badge = p_badge.add_run("AGRITECH AI HACKATHON 2026 · AQMOLA HUB · DIGITAL AQMOLA")
    r_badge.font.name = 'Arial'
    r_badge.font.size = Pt(9.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = C_GREEN

    p_title = c.add_paragraph()
    r_title = p_title.add_run("DalaSat — Инструкция по локальному запуску прототипа")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = C_DARK

    p_sub = c.add_paragraph()
    r_sub = p_sub.add_run("Руководство для экспертной комиссии по проверке воспроизводимости решения на чистой машине")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = C_GRAY

    doc.add_paragraph()

    # ─── 2. ИНФОРМАЦИЯ О ПРОЕКТЕ И КОМАНДЕ ─────────────────────────────────────
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Общие сведения о проекте и команде")
    r_h1.font.name = 'Arial'
    r_h1.font.color.rgb = C_GREEN

    t_info = doc.add_table(rows=5, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_info.autofit = False

    info_data = [
        ("Название проекта", "DalaSat (Спутниковый экспресс-мониторинг полей и ИИ-агроном)"),
        ("Выбранный трек", "Трек 1. GIS и дистанционное зондирование Земли"),
        ("Хакатон и организатор", "AI-хакатон «AgriTech AI Hackathon» · Акмолинский филиал АКФ «Астана Хаб» (Aqmola Hub) при поддержке Управления цифровизации и архивов Акмолинской области (Digital Aqmola)"),
        ("Состав команды", "• Адильхан Ануар — Founder · IT / Design\n• Алиев Исмайл — Design · Marketing\n• Досай Али — Cofounder · Speaker"),
        ("Ожидаемый результат по ТЗ", "Telegram-бот со слоями спутниковых индексов Sentinel-2 (NDVI, NDMI), историей поля, детекцией объектов инфраструктуры и выгрузкой отчёта в PDF")
    ]

    for i, (k, v) in enumerate(info_data):
        row = t_info.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.6)
        set_cell_margins(c0, 100, 100, 120, 120)
        set_cell_margins(c1, 100, 100, 120, 120)
        set_cell_background(c0, "F9F8F5")
        set_cell_background(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.name = 'Arial'
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = C_DARK

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = 'Arial'
        r1.font.size = Pt(10)
        r1.font.color.rgb = C_DARK

    doc.add_paragraph()

    # ─── 3. СИСТЕМНЫЕ ТРЕБОВАНИЯ ──────────────────────────────────────────────
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Системные требования")
    r_h2.font.name = 'Arial'
    r_h2.font.color.rgb = C_GREEN

    p_req = doc.add_paragraph()
    p_req.paragraph_format.line_spacing = 1.2
    reqs = [
        ("Операционная система: ", "Windows 10/11, macOS (Intel / Apple Silicon) или Linux (Ubuntu 20.04+)."),
        ("Интерпретатор Python: ", "Python 3.10, 3.11 или 3.12 (рекомендуется Python 3.11/3.12)."),
        ("Оперативная память: ", "От 2 ГБ RAM (решение крайне легковесно, так как использует технологию windowed read и обрабатывает фрагменты по 200 КБ вместо гигабайтных снимков)."),
        ("Свободное место на диске: ", "Около 300–400 МБ для виртуального окружения и библиотек."),
        ("Сетевое подключение: ", "Стабильный доступ к интернету для обращения к STAC API Microsoft Planetary Computer, Telegram API и Google Gemini API.")
    ]
    for bold_prefix, text in reqs:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(3)
        rb = p_item.add_run(bold_prefix)
        rb.font.name = 'Arial'
        rb.font.bold = True
        rb.font.size = Pt(10.5)
        rb.font.color.rgb = C_DARK
        rt = p_item.add_run(text)
        rt.font.name = 'Arial'
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = C_DARK

    doc.add_paragraph()

    # ─── 4. ПОШАГОВЫЙ ЛОКАЛЬНЫЙ ЗАПУСК ─────────────────────────────────────────
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. Пошаговая инструкция запуска «с нуля»")
    r_h3.font.name = 'Arial'
    r_h3.font.color.rgb = C_GREEN

    # Шаг 1
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(8)
    p_s1.paragraph_format.space_after = Pt(4)
    r_s1_num = p_s1.add_run("Шаг 1. Клонирование репозитория или переход в папку проекта\n")
    r_s1_num.font.name = 'Arial'
    r_s1_num.font.bold = True
    r_s1_num.font.size = Pt(12)
    r_s1_num.font.color.rgb = C_GREEN

    p_cmd1 = doc.add_paragraph()
    p_cmd1.paragraph_format.left_indent = Inches(0.2)
    r_cmd1 = p_cmd1.add_run("git clone https://github.com/your-team/agro_sat_bot.git\ncd agro_sat_bot")
    r_cmd1.font.name = 'Consolas'
    r_cmd1.font.size = Pt(9.5)
    r_cmd1.font.bold = True

    # Шаг 2
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_before = Pt(8)
    p_s2.paragraph_format.space_after = Pt(4)
    r_s2_num = p_s2.add_run("Шаг 2. Создание и активация виртуального окружения (Virtualenv)\n")
    r_s2_num.font.name = 'Arial'
    r_s2_num.font.bold = True
    r_s2_num.font.size = Pt(12)
    r_s2_num.font.color.rgb = C_GREEN

    r_s2_desc = p_s2.add_run("Для изоляции зависимостей создайте чистое виртуальное окружение:")
    r_s2_desc.font.name = 'Arial'
    r_s2_desc.font.size = Pt(10.5)

    p_cmd2 = doc.add_paragraph()
    p_cmd2.paragraph_format.left_indent = Inches(0.2)
    r_cmd2 = p_cmd2.add_run("# 1. Создание окружения:\npython -m venv .venv\n\n# 2. Активация на Windows (PowerShell / CMD):\n.venv\\Scripts\\activate\n\n# 2. Активация на macOS / Linux:\nsource .venv/bin/activate")
    r_cmd2.font.name = 'Consolas'
    r_cmd2.font.size = Pt(9.5)
    r_cmd2.font.bold = True

    # Шаг 3
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_before = Pt(8)
    p_s3.paragraph_format.space_after = Pt(4)
    r_s3_num = p_s3.add_run("Шаг 3. Установка библиотек из requirements.txt\n")
    r_s3_num.font.name = 'Arial'
    r_s3_num.font.bold = True
    r_s3_num.font.size = Pt(12)
    r_s3_num.font.color.rgb = C_GREEN

    p_cmd3 = doc.add_paragraph()
    p_cmd3.paragraph_format.left_indent = Inches(0.2)
    r_cmd3 = p_cmd3.add_run("pip install --upgrade pip\npip install -r requirements.txt")
    r_cmd3.font.name = 'Consolas'
    r_cmd3.font.size = Pt(9.5)
    r_cmd3.font.bold = True

    # Примечание по Rasterio на Windows
    t_note = doc.add_table(rows=1, cols=1)
    c_note = t_note.cell(0, 0)
    set_cell_background(c_note, HEX_BEIGE)
    set_cell_margins(c_note, 100, 100, 150, 150)
    p_n = c_note.paragraphs[0]
    r_nb = p_n.add_run("Важное примечание для пользователей Windows: ")
    r_nb.font.name = 'Arial'
    r_nb.font.bold = True
    r_nb.font.size = Pt(9.5)
    r_nt = p_n.add_run("Все основные библиотеки (aiogram, pystac-client, planetary-computer, rasterio, reportlab, shapely, google-genai) ставятся в виде готовых предкомпилированных бинарных пакетов (wheels). Если в вашей системе Windows отсутствует компилятор C++, rasterio без проблем ставится стандартной командой pip install rasterio.")
    r_nt.font.name = 'Arial'
    r_nt.font.size = Pt(9.5)

    # Шаг 4
    p_s4 = doc.add_paragraph()
    p_s4.paragraph_format.space_before = Pt(10)
    p_s4.paragraph_format.space_after = Pt(4)
    r_s4_num = p_s4.add_run("Шаг 4. Настройка файла переменных окружения (.env)\n")
    r_s4_num.font.name = 'Arial'
    r_s4_num.font.bold = True
    r_s4_num.font.size = Pt(12)
    r_s4_num.font.color.rgb = C_GREEN

    r_s4_desc = p_s4.add_run("В корне проекта создайте файл .env (или отредактируйте существующий) со следующими двумя ключами:")
    r_s4_desc.font.name = 'Arial'
    r_s4_desc.font.size = Pt(10.5)

    p_cmd4 = doc.add_paragraph()
    p_cmd4.paragraph_format.left_indent = Inches(0.2)
    r_cmd4 = p_cmd4.add_run("TELEGRAM_BOT_TOKEN=8800586784:AAH...ваш_токен_бота...\nGEMINI_API_KEY=AIzaSy...ваш_api_ключ_gemini...")
    r_cmd4.font.name = 'Consolas'
    r_cmd4.font.size = Pt(9.5)
    r_cmd4.font.bold = True

    p_s4_exp = doc.add_paragraph()
    p_s4_exp.paragraph_format.line_spacing = 1.15
    r_s4_exp = p_s4_exp.add_run(
        "1. Telegram Bot Token: Создаётся за 30 секунд в Telegram через официального бота @BotFather командой /newbot.\n"
        "2. Gemini API Key: Бесплатно генерируется в Google AI Studio (aistudio.google.com) по кнопке «Get API Key». В проекте используется отказоустойчивый нативный Interactions API с каскадной ротацией моделей (gemini-3.5-flash-lite), поэтому даже при отсутствии ключа система автоматически активирует резервный агрономический модуль.\n"
        "3. Доступ к спутникам Sentinel-2: Не требует платных подписок или ключей! Запросы выполняются напрямую к бесплатному открытому каталогу Microsoft Planetary Computer STAC API."
    )
    r_s4_exp.font.name = 'Arial'
    r_s4_exp.font.size = Pt(10)

    # Шаг 5
    p_s5 = doc.add_paragraph()
    p_s5.paragraph_format.space_before = Pt(8)
    p_s5.paragraph_format.space_after = Pt(4)
    r_s5_num = p_s5.add_run("Шаг 5. Запуск Telegram-бота\n")
    r_s5_num.font.name = 'Arial'
    r_s5_num.font.bold = True
    r_s5_num.font.size = Pt(12)
    r_s5_num.font.color.rgb = C_GREEN

    p_cmd5 = doc.add_paragraph()
    p_cmd5.paragraph_format.left_indent = Inches(0.2)
    r_cmd5 = p_cmd5.add_run("python bot.py")
    r_cmd5.font.name = 'Consolas'
    r_cmd5.font.size = Pt(10)
    r_cmd5.font.bold = True

    p_s5_ok = doc.add_paragraph()
    r_s5_ok = p_s5_ok.add_run("При успешном запуске в терминале отобразится лог:\n")
    r_s5_ok.font.name = 'Arial'
    r_s5_ok.font.size = Pt(10)

    p_cmd5_log = doc.add_paragraph()
    p_cmd5_log.paragraph_format.left_indent = Inches(0.2)
    r_cmd5_log = p_cmd5_log.add_run("INFO | __main__ | 🌾 DalaSat Bot успешно запущен и готов к приёму координат!")
    r_cmd5_log.font.name = 'Consolas'
    r_cmd5_log.font.size = Pt(9.5)
    r_cmd5_log.font.color.rgb = C_GREEN

    doc.add_paragraph()

    # ─── 5. ТЕСТОВЫЙ СЦЕНАРИЙ ДЛЯ ЭКСПЕРТА ────────────────────────────────────
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Контрольный сквозной сценарий тестирования экспертом")
    r_h4.font.name = 'Arial'
    r_h4.font.color.rgb = C_GREEN

    p_sc = doc.add_paragraph()
    p_sc.paragraph_format.line_spacing = 1.25
    r_sc = p_sc.add_run(
        "Члены экспертной комиссии могут повторить проверку за 1 минуту по следующему шаговому чек-листу:\n\n"
        "1. Откройте запущенного бота в Telegram и отправьте команду /start.\n"
        "2. Выберите язык интерфейса: Русский (RU) или Казахский (KZ).\n"
        "3. Выберите масштаб поля кнопкой в меню: например, «🌾 Стандартный (2×2 км, ~400 га)».\n"
        "4. Отправьте координаты реального зернового поля в Акмолинской области. Вы можете нажать кнопку «📍 Отправить геопозицию» со смартфона или просто скопировать и отправить текст:\n"
    )
    r_sc.font.name = 'Arial'
    r_sc.font.size = Pt(10)

    p_coords = doc.add_paragraph()
    p_coords.paragraph_format.left_indent = Inches(0.2)
    r_coords = p_coords.add_run("53.2514, 69.1845")
    r_coords.font.name = 'Consolas'
    r_coords.font.size = Pt(11)
    r_coords.font.bold = True
    r_coords.font.color.rgb = C_GREEN

    p_sc_cont = doc.add_paragraph()
    p_sc_cont.paragraph_format.line_spacing = 1.25
    r_sc_cont = p_sc_cont.add_run(
        "(Это реальный тестовый массив пашни в Зерендинском районе Акмолинской области).\n\n"
        "5. Ожидание результата (8–10 секунд): бот выполнит запрос к Sentinel-2, нормализует каналы Red, NIR, SWIR, рассчитает матрицы NDVI и NDMI, выявит строения через NDBI + BFS и сформирует экспертный вердикт Gemini AI.\n"
        "6. Проверка ответа: в чат поступит цветная тепловая карта поля высокой четкости со шкалой биомассы, структурированные показатели (готовность почвы к севу в %, вегетация, влажность, постройки) и зональные агрономические советы.\n"
        "7. Проверка PDF: нажмите кнопку «📄 Скачать PDF-паспорт». Бот мгновенно пришлёт сформированный одностраничный официальный PDF-паспорт с картой и таблицами для субсидирования."
    )
    r_sc_cont.font.name = 'Arial'
    r_sc_cont.font.size = Pt(10)

    doc.add_paragraph()

    # ─── 6. АРХИТЕКТУРА И СООТВЕТСТВИЕ ТЗ ─────────────────────────────────────
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. Архитектура файлов и соответствие ТЗ")
    r_h5.font.name = 'Arial'
    r_h5.font.color.rgb = C_GREEN

    t_files = doc.add_table(rows=8, cols=2)
    t_files.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_files.autofit = False

    files_data = [
        ("bot.py", "Главный файл Telegram-бота на Aiogram 3. Обработка команд, геолокаций, масштабирования, переключения языков (KZ/RU)."),
        ("config.py", "Конфигурация системы: токены, координаты дельты для масштабов (1x1, 2x2, 4x4 км), пороги спектральных индексов."),
        ("services/satellite_service.py", "Интеграция с Microsoft Planetary Computer STAC API. Скачивание COG Sentinel-2 L2A оконным чтением."),
        ("services/gis_service.py", "Расчёт индексов биомассы NDVI, влажности NDMI, нормализация растров в NumPy и генерация цветовых тепловых карт."),
        ("services/building_detector.py", "Детекция инфраструктуры (ангары, склады, станы) на основе индекса NDBI и алгоритма связности BFS."),
        ("services/ai_advisor.py", "Интеллектуальный агроном на базе Google Gemini Interactions API с зональной базой знаний Акмолинской области."),
        ("services/pdf_generator.py", "Генератор официального PDF-паспорта поля с векторной версткой и таблицами на ReportLab."),
        ("build_html_presentation.py", "Генератор официальной 10-слайдовой презентации хакатона (16:9 PDF через Chromium engine).")
    ]

    for i, (fname, fdesc) in enumerate(files_data):
        row = t_files.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.6)
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)
        set_cell_background(c0, "F9F8F5")
        set_cell_background(c1, "FFFFFF")

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(fname)
        r0.font.name = 'Consolas'
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = C_GREEN

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(fdesc)
        r1.font.name = 'Arial'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = C_DARK

    doc.add_paragraph()

    # ─── 7. ВОЗМОЖНЫЕ ВОПРОСЫ И УСТРАНЕНИЕ ОШИБОК ──────────────────────────────
    h6 = doc.add_heading(level=1)
    r_h6 = h6.add_run("6. Диагностика и устранение неполадок (Troubleshooting)")
    r_h6.font.name = 'Arial'
    r_h6.font.color.rgb = C_GREEN

    faq = [
        ("Ошибка «TelegramUnauthorizedError: Unauthorized»:", "Проверьте токен в файле .env. Убедитесь, что токен скопирован полностью из @BotFather без лишних пробелов."),
        ("Ошибка «Rate limit / ResourceExhausted» в Gemini API:", "В сервисе ai_advisor.py уже реализован автоматический каскадный Fallback. Даже если бесплатный лимит Google AI Studio исчерпан, бот не падает, а выдает проверенные зональные рекомендации из встроенного экспертного агрономического справочника."),
        ("Ошибка при открытии COG ссылки Planetary Computer:", "Планетарный компьютер Microsoft иногда обновляет SAS-токены. Сервис satellite_service.py автоматически подписывает каждую ссылку через planetary_computer.sign_url(). Убедитесь, что интернет-соединение активно.")
    ]

    for q, a in faq:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(2)
        rq = p_q.add_run(q)
        rq.font.name = 'Arial'
        rq.font.bold = True
        rq.font.size = Pt(10)
        rq.font.color.rgb = C_GREEN

        p_a = doc.add_paragraph()
        p_a.paragraph_format.left_indent = Inches(0.15)
        p_a.paragraph_format.space_after = Pt(6)
        ra = p_a.add_run(a)
        ra.font.name = 'Arial'
        ra.font.size = Pt(9.5)
        ra.font.color.rgb = C_DARK

    # ─── СОХРАНЕНИЕ ───────────────────────────────────────────────────────────
    project_doc = os.path.join(DIR, "DalaSat_Инструкция_по_запуску.docx")
    docs_folder = os.path.expanduser("~/Documents")
    user_doc = os.path.join(docs_folder, "DalaSat_Инструкция_по_запуску.docx")

    doc.save(project_doc)
    doc.save(user_doc)
    print(f"Word manual created successfully:")
    print(f"1. Project: {project_doc}")
    print(f"2. Documents: {user_doc}")

if __name__ == "__main__":
    create_manual()
